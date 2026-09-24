import unittest
from io import BytesIO
from unittest.mock import patch

from openpyxl import load_workbook

import app as app_module
import routes.users_roles as users_roles
from rbac_catalog import PERMISOS_ROLES_SISTEMA
from services.tenant_backup import (
    TABLAS_PERMITIDAS, generar_backup_excel, nombre_archivo_backup,
)


class TenantBackupTests(unittest.TestCase):
    def test_excel_has_one_sheet_per_module_and_omits_secrets(self):
        def consultar(query, params=None, fetch='one'):
            if query.startswith('SHOW COLUMNS FROM `usuarios`'):
                return [
                    {'Field': 'id'},
                    {'Field': 'tenant_id'},
                    {'Field': 'nombre'},
                    {'Field': 'password_hash'},
                    {'Field': 'reset_token'},
                ]
            if query.startswith('SHOW COLUMNS FROM `empresas`'):
                return [{'Field': 'id'}, {'Field': 'nombre'}]
            if query.startswith('SHOW COLUMNS'):
                return [{'Field': 'id'}]
            if query.startswith('SELECT') and 'FROM `usuarios`' in query:
                self.assertEqual(params, (9,))
                self.assertNotIn('password_hash', query)
                self.assertNotIn('reset_token', query)
                return [{'id': 4, 'tenant_id': 9, 'nombre': 'Ana'}]
            if query.startswith('SELECT') and 'FROM `empresas`' in query:
                self.assertEqual(params, (9,))
                return [{'id': 9, 'nombre': 'Centro Norte'}]
            return []

        libro = load_workbook(generar_backup_excel(9, consultar=consultar))
        self.assertIn('Resumen', libro.sheetnames)
        self.assertIn('Pacientes', libro.sheetnames)
        self.assertIn('Facturas', libro.sheetnames)
        self.assertIn('Usuarios', libro.sheetnames)
        self.assertNotIn('Roles', libro.sheetnames)
        self.assertNotIn('Citas', libro.sheetnames)
        self.assertNotIn('Turnos', libro.sheetnames)
        self.assertNotIn('Mensajeria', libro.sheetnames)
        usuarios = libro['Usuarios']
        encabezados = [celda.value for celda in next(usuarios.iter_rows(min_row=2, max_row=2))]
        self.assertIn('nombre', encabezados)
        self.assertNotIn('password_hash', encabezados)
        self.assertNotIn('reset_token', encabezados)
        excluidas = {
            'auditoria_licencias_medicas', 'auditoria_historia_clinica',
            'evoluciones_clinicas', 'citas_medicas', 'roles', 'usuario_roles',
            'usuario_medico', 'pacientes_pendientes', 'centros_medicos',
            'medico_centro', 'codigo_ars', 'servicios', 'ncf',
            'conversaciones_internas', 'mensajes_internos',
            'turnos_atencion', 'turnos_eventos', 'pantallas_turnos',
            'secuencias_turnos',
        }
        self.assertTrue(excluidas.isdisjoint(TABLAS_PERMITIDAS))

    def test_backup_filename_uses_company_name(self):
        nombre = nombre_archivo_backup(
            3,
            consultar=lambda *args, **kwargs: {'nombre': 'Centro Norte'},
        )
        self.assertTrue(nombre.startswith('ClinicRD_backup_Centro_Norte_'))
        self.assertTrue(nombre.endswith('.xlsx'))

    def test_backup_events_report_module_progress(self):
        from services.tenant_backup import generar_backup_eventos

        def consultar(query, params=None, fetch='one'):
            if query.startswith('SHOW COLUMNS FROM `empresas`'):
                return [{'Field': 'id'}, {'Field': 'nombre'}]
            if query.startswith('SHOW COLUMNS'):
                return [{'Field': 'id'}]
            if 'FROM empresas WHERE id' in query:
                return {'nombre': 'Centro Norte'}
            return []

        eventos = list(generar_backup_eventos(9, consultar=consultar))
        self.assertGreaterEqual(len(eventos), 3)
        self.assertEqual(eventos[0]['pct'], 4)
        self.assertIn('contenido', eventos[-1])
        self.assertEqual(eventos[-1]['pct'], 100)
        self.assertTrue(eventos[-1]['nombre'].endswith('.xlsx'))
        self.assertTrue(eventos[-1]['contenido'].startswith(b'PK'))

    def test_only_administrator_has_backup_permission(self):
        self.assertIn('configuracion.backup', PERMISOS_ROLES_SISTEMA['Administrador'])
        for rol in ('Nivel 2', 'Registro de Facturas', 'Oficial de servicios', 'Médico'):
            self.assertNotIn('configuracion.backup', PERMISOS_ROLES_SISTEMA[rol])

    def test_backup_route_requires_tenant(self):
        handler = users_roles.perfil_descargar_backup
        while hasattr(handler, '__wrapped__'):
            handler = handler.__wrapped__
        with app_module.app.test_request_context(
            '/perfil/configuracion/backup'
        ), patch.object(
            users_roles, 'get_current_tenant_id', return_value=None
        ), patch.object(
            users_roles, 'flash'
        ), patch.object(
            users_roles, 'redirect', return_value='redir'
        ) as redirect:
            respuesta = handler()
        self.assertEqual(respuesta, 'redir')
        redirect.assert_called()
