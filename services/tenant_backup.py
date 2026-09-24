"""Copia Excel de los datos de un tenant, una hoja por módulo."""

import logging
import re
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.database import execute_query

logger = logging.getLogger(__name__)

_IDENTIFICADOR = re.compile(r'^[A-Za-z_][A-Za-z0-9_]*$')
_SECRETOS = (
    'password', 'passwd', 'token', 'secret', 'secreto', 'hash',
    'cifrado', 'private', 'p12',
)

MODULOS = (
    ('Empresa', (('empresas', 'id'),)),
    ('Usuarios', (('usuarios', 'tenant_id'),)),
    ('Pacientes', (('pacientes', 'tenant_id'),)),
    ('Medicos', (('medicos', 'tenant_id'),)),
    ('Catalogos', (('ars', 'tenant_id'),)),
    ('Historia clinica', (('consultas_clinicas', 'tenant_id'),)),
    ('Recetas', (
        ('recetas_medicas', 'tenant_id'),
        ('receta_medicamentos', 'tenant_id'),
    )),
    ('Licencias', (
        ('tipos_licencia_medica', 'tenant_id'),
        ('licencias_medicas', 'tenant_id'),
    )),
    ('Emergencia', (('historias_emergencia', 'tenant_id'),)),
    ('Enfermeria', (('hojas_enfermeria', 'tenant_id'),)),
    ('Facturas', (
        ('facturas', 'tenant_id'),
        ('factura_detalles', 'tenant_id'),
        ('facturas_ecf', 'tenant_id'),
    )),
    ('Pagos', (
        ('pagos', 'tenant_id'),
        ('pago_facturas', 'tenant_id'),
    )),
    ('Reclamaciones', (('reclamaciones', 'tenant_id'),)),
    ('Facturacion electronica', (
        ('ecf_configuraciones', 'tenant_id'),
        ('ecf_secuencias', 'tenant_id'),
        ('ecf_eventos', 'tenant_id'),
        ('ecf_outbox', 'tenant_id'),
    )),
)

TABLAS_PERMITIDAS = frozenset(
    tabla for _, tablas in MODULOS for tabla, _filtro in tablas
)


def _es_identificador(nombre):
    return bool(nombre and _IDENTIFICADOR.match(str(nombre)))


def _es_columna_secreta(nombre):
    clave = str(nombre or '').casefold()
    return any(parte in clave for parte in _SECRETOS)


def _valor_celda(valor):
    if valor is None:
        return ''
    if isinstance(valor, (datetime, date)):
        return valor.isoformat(sep=' ', timespec='seconds') if isinstance(valor, datetime) else valor.isoformat()
    if isinstance(valor, Decimal):
        return float(valor)
    if isinstance(valor, (bytes, bytearray, memoryview)):
        return '[binario]'
    if isinstance(valor, (dict, list, tuple)):
        return str(valor)
    return valor


def _columnas_exportables(tabla, consultar=execute_query):
    if tabla not in TABLAS_PERMITIDAS or not _es_identificador(tabla):
        return []
    try:
        filas = consultar(f'SHOW COLUMNS FROM `{tabla}`', fetch='all') or []
    except Exception as exc:
        logger.warning('No se pudo leer columnas de %s: %s', tabla, exc)
        return []
    columnas = []
    for fila in filas:
        campo = fila.get('Field') or fila.get('field')
        if _es_identificador(campo) and not _es_columna_secreta(campo):
            columnas.append(campo)
    return columnas


def _filas_tabla(tabla, filtro, tenant_id, consultar=execute_query):
    columnas = _columnas_exportables(tabla, consultar)
    if not columnas:
        return [], []
    if filtro not in {'id', 'tenant_id'}:
        return columnas, []
    seleccion = ', '.join(f'`{columna}`' for columna in columnas)
    try:
        filas = consultar(
            f'SELECT {seleccion} FROM `{tabla}` WHERE `{filtro}`=%s',
            (tenant_id,),
            fetch='all',
        ) or []
    except Exception as exc:
        logger.warning('No se pudo leer %s del tenant %s: %s', tabla, tenant_id, exc)
        return columnas, []
    return columnas, filas


def generar_backup_eventos(tenant_id, consultar=execute_query):
    """Ir armando el Excel y avisar el avance de cada módulo."""
    if not tenant_id:
        raise ValueError('Se requiere la empresa para generar el backup')
    yield {'pct': 4, 'label': 'Preparando copia de seguridad'}
    wb = Workbook()
    resumen = wb.active
    resumen.title = 'Resumen'
    header_fill = PatternFill(start_color='0F766E', end_color='0F766E', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')
    titulo_fill = PatternFill(start_color='ECFDF5', end_color='ECFDF5', fill_type='solid')
    titulo_font = Font(bold=True, color='0F766E')

    for col, titulo in enumerate(('Modulo', 'Tabla', 'Registros'), 1):
        celda = resumen.cell(row=1, column=col, value=titulo)
        celda.fill = header_fill
        celda.font = header_font
        celda.alignment = Alignment(horizontal='center')

    fila_resumen = 2
    total = len(MODULOS)
    for indice, (nombre_modulo, tablas) in enumerate(MODULOS, 1):
        hoja = wb.create_sheet(nombre_modulo[:31])
        fila = 1
        for tabla, filtro in tablas:
            columnas, registros = _filas_tabla(tabla, filtro, tenant_id, consultar)
            hoja.cell(row=fila, column=1, value=tabla).font = titulo_font
            hoja.cell(row=fila, column=1).fill = titulo_fill
            fila += 1
            if columnas:
                for col, columna in enumerate(columnas, 1):
                    celda = hoja.cell(row=fila, column=col, value=columna)
                    celda.fill = header_fill
                    celda.font = header_font
                fila += 1
                for registro in registros:
                    for col, columna in enumerate(columnas, 1):
                        hoja.cell(
                            row=fila,
                            column=col,
                            value=_valor_celda(registro.get(columna)),
                        )
                    fila += 1
            else:
                hoja.cell(row=fila, column=1, value='Sin datos o tabla no disponible')
                fila += 1
            fila += 1
            resumen.cell(row=fila_resumen, column=1, value=nombre_modulo)
            resumen.cell(row=fila_resumen, column=2, value=tabla)
            resumen.cell(row=fila_resumen, column=3, value=len(registros))
            fila_resumen += 1
        for col in range(1, 8):
            hoja.column_dimensions[get_column_letter(col)].width = 22
        yield {
            'pct': 4 + int(indice * 86 / total),
            'label': f'Exportando {nombre_modulo}',
        }

    for col, ancho in enumerate((22, 32, 14), 1):
        resumen.column_dimensions[get_column_letter(col)].width = ancho

    yield {'pct': 94, 'label': 'Cerrando el libro Excel'}
    salida = BytesIO()
    wb.save(salida)
    yield {
        'pct': 100,
        'label': 'Descarga lista',
        'nombre': nombre_archivo_backup(tenant_id, consultar),
        'contenido': salida.getvalue(),
    }


def generar_backup_excel(tenant_id, consultar=execute_query):
    """Armar el libro Excel con los módulos de la empresa."""
    archivo = None
    for evento in generar_backup_eventos(tenant_id, consultar):
        contenido = evento.get('contenido')
        if contenido is not None:
            archivo = BytesIO(contenido)
            archivo.seek(0)
    return archivo


def nombre_archivo_backup(tenant_id, consultar=execute_query):
    empresa = consultar(
        'SELECT nombre FROM empresas WHERE id=%s',
        (tenant_id,),
    ) or {}
    etiqueta = re.sub(r'[^A-Za-z0-9_-]+', '_', str(empresa.get('nombre') or 'empresa'))[:40]
    fecha = datetime.now().strftime('%Y%m%d_%H%M%S')
    return f'ClinicRD_backup_{etiqueta}_{fecha}.xlsx'
